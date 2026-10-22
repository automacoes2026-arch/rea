#!/usr/bin/env python3
"""Contagens de uma base de prospecção PJ (regra 7.4 do Sistema Mestre).

Determinístico e local: lê um CSV ou XLSX, valida CNPJ pelos dígitos verificadores,
deduplica e emite apenas contagens agregadas. Nenhum CNPJ, nome ou contato é
gravado na saída.

Uso:
    python3 contagens_base.py BASE.csv --config mapeamento.json [--saida contagens.json]
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from pathlib import Path

VERSAO = "1.0.0"

# Contagens por entidade, na ordem da regra 7.4. Cada uma é uma coluna opcional do
# mapeamento; uma entidade conta quando qualquer linha dela tem valor verdadeiro.
FLAGS = (
    "no_segmento",
    "enriquecida",
    "canal_corporativo",
    "priorizada",
    "abordada",
    "oportunidade_aberta",
)

VERDADEIROS_PADRAO = ("1", "s", "sim", "true", "verdadeiro", "x", "yes")


def cnpj_digitos(valor: object) -> str:
    return "".join(c for c in str(valor or "") if c.isdigit())


def cnpj_valido(digitos: str) -> bool:
    if len(digitos) != 14 or digitos == digitos[0] * 14:
        return False

    def dv(base: str, pesos: tuple[int, ...]) -> str:
        resto = sum(int(d) * p for d, p in zip(base, pesos)) % 11
        return "0" if resto < 2 else str(11 - resto)

    p1 = (5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2)
    p2 = (6,) + p1
    d1 = dv(digitos[:12], p1)
    return digitos[12] == d1 and digitos[13] == dv(digitos[:12] + d1, p2)


def ler_linhas(caminho: Path) -> list[dict[str, object]]:
    if caminho.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook  # dependência só para XLSX

        planilha = load_workbook(caminho, read_only=True, data_only=True).active
        linhas = planilha.iter_rows(values_only=True)
        cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas, ())]
        return [dict(zip(cabecalho, linha)) for linha in linhas if any(v is not None for v in linha)]
    with caminho.open(newline="", encoding="utf-8-sig") as arquivo:
        amostra = arquivo.read(4096)
        arquivo.seek(0)
        dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
        return list(csv.DictReader(arquivo, dialect=dialeto))


def contar(linhas: list[dict[str, object]], config: dict) -> dict:
    coluna_cnpj = config["cnpj"]
    colunas = config.get("colunas", {})
    verdadeiros = {v.lower() for v in config.get("verdadeiros", VERDADEIROS_PADRAO)}
    chave = config.get("chave", "raiz")
    if chave not in ("raiz", "completo"):
        raise ValueError("chave deve ser 'raiz' ou 'completo'")
    desconhecidas = set(colunas) - set(FLAGS)
    if desconhecidas:
        raise ValueError(f"contagens desconhecidas no mapeamento: {sorted(desconhecidas)}")
    if linhas and coluna_cnpj not in linhas[0]:
        raise ValueError(f"coluna de CNPJ '{coluna_cnpj}' não existe no arquivo")
    ausentes = sorted(c for c in colunas.values() if linhas and c not in linhas[0])
    if ausentes:
        raise ValueError(f"colunas mapeadas não existem no arquivo: {ausentes}")

    invalidos = 0
    entidades: dict[str, set[str]] = {}
    completos: set[str] = set()
    for linha in linhas:
        digitos = cnpj_digitos(linha.get(coluna_cnpj))
        if not cnpj_valido(digitos):
            invalidos += 1
            continue
        completos.add(digitos)
        marcas = entidades.setdefault(digitos[:8] if chave == "raiz" else digitos, set())
        for flag, coluna in colunas.items():
            if str(linha.get(coluna) or "").strip().lower() in verdadeiros:
                marcas.add(flag)

    validos = len(linhas) - invalidos
    resultado = {
        "registros_totais": len(linhas),
        "registros_cnpj_invalido": invalidos,
        "cadastros_validos": validos,
        "cnpjs_unicos": len(completos),
        "entidades_unicas": len(entidades),
        "registros_duplicados": validos - len(entidades),
    }
    for flag in FLAGS:
        resultado[flag] = (
            sum(flag in marcas for marcas in entidades.values()) if flag in colunas else None
        )
    return resultado


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("base", type=Path)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--saida", type=Path)
    args = parser.parse_args(argv)

    config = json.loads(args.config.read_text(encoding="utf-8"))
    try:
        contagens = contar(ler_linhas(args.base), config)
    except (ValueError, KeyError) as erro:
        print(f"erro: {erro}", file=sys.stderr)
        return 2

    relatorio = {
        "versao_script": VERSAO,
        "sha256_entrada": hashlib.sha256(args.base.read_bytes()).hexdigest(),
        "chave_entidade": config.get("chave", "raiz"),
        "contagens": contagens,
        "nao_mapeadas": [f for f in FLAGS if contagens[f] is None],
    }
    texto = json.dumps(relatorio, ensure_ascii=False, indent=2)
    if args.saida:
        args.saida.write_text(texto + "\n", encoding="utf-8")
    print(texto)
    return 0


if __name__ == "__main__":
    sys.exit(main())
